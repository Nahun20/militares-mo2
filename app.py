import hashlib
import hmac
import io
import json
import os
import time

import pandas as pd
import streamlit as st
from google import genai
from google.genai import types
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps
from pydantic import BaseModel

# Modelo con nivel gratuito. Si deja de funcionar, prueba "gemini-3.5-flash-lite".
MODELO = "gemini-3.8-flash"

st.set_page_config(page_title="Control de conteo", page_icon="📋", layout="wide")
st.title("📋 Control de conteo: foto → Excel")


# ---------- cliente de la API ----------
def get_client():
    key = os.getenv("GEMINI_API_KEY")
    if not key:
        try:
            key = st.secrets["GEMINI_API_KEY"]
        except Exception:
            key = None
    if not key:
        st.error(
            "Falta la API key. Crea el archivo `.streamlit/secrets.toml` con la línea "
            '`GEMINI_API_KEY = "tu_clave"` y reinicia la app.'
        )
        st.stop()
    return genai.Client(api_key=key)


# ---------- validación de EAN-13 ----------
def ean_valido(code: str) -> bool:
    code = str(code or "")
    if len(code) != 13 or not code.isdigit():
        return False
    suma = sum(int(d) * (3 if i % 2 else 1) for i, d in enumerate(code[:12]))
    return (10 - suma % 10) % 10 == int(code[12])


# ---------- imagen ----------
def preparar_imagen(img_bytes: bytes) -> bytes:
    """Corrige la rotación de WhatsApp/celular y reduce el tamaño para la API."""
    img = Image.open(io.BytesIO(img_bytes))
    img = ImageOps.exif_transpose(img).convert("RGB")
    img.thumbnail((2000, 2000))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=90)
    return buf.getvalue()


# ---------- lectura con Gemini ----------
class Fila(BaseModel):
    caja: int
    ean: str
    descripcion: str
    piezas: int
    nombre: str
    dudosa: bool
    nota: str


class Hoja(BaseModel):
    fecha: str
    sede: str
    tarima: str
    supervisor: str
    filas: list[Fila]


PROMPT = """Esta es una hoja "Control de conteo de mercancía por caja" llenada a mano.
Transcribe TODAS las filas que tengan datos, de arriba hacia abajo, y devuelve el JSON pedido.

Reglas:
- Transcribe lo que está escrito. No inventes ni "corrijas" números.
- EAN: solo dígitos (normalmente 13 y empiezan con 750). Si un carácter no se entiende,
  escríbelo como "?". No completes dígitos que no ves.
- Las comillas (", '', ") significan "igual que la fila de arriba". Si el EAN trae comillas,
  repite el EAN de arriba. En la descripción, repite el texto de arriba y agrega lo que sí
  esté escrito (ej. arriba "Tinta multiusos Parvula" y abajo '" " Fugaz' -> "Tinta multiusos Fugaz").
- Caja: el número tal como está escrito (puede repetirse en varias filas).
- Piezas: número entero.
- Nombre: quien realizó el conteo, tal como está escrito.
- Ignora encabezados, filas vacías y el total.
- dudosa=true cuando algún dato de la fila sea difícil de leer; en "nota" explica qué y por qué.
  Si la fila es clara, dudosa=false y nota es cadena vacía.
- Encabezado de la hoja (fuera de la tabla), tal como está escrito:
  fecha = lo escrito junto a "FECHA:"; sede = junto a "SEDE:"; tarima = junto a "Tarima";
  supervisor = nombre junto a "SUPERVISOR:" (suele estar al final de la hoja, puede ser una firma:
  si no se entiende el nombre, escribe lo que alcances a leer).
  Si alguno de esos campos está vacío, devuelve cadena vacía. No inventes valores."""


@st.cache_data(show_spinner=False)
def leer_hoja(img_bytes: bytes, _client):
    jpg = preparar_imagen(img_bytes)
    config = types.GenerateContentConfig(
        response_mime_type="application/json",
        response_schema=Hoja,
        temperature=0,
    )
    ultimo_error = None
    for intento in range(3):
        try:
            r = _client.models.generate_content(
                model=MODELO,
                contents=[types.Part.from_bytes(data=jpg, mime_type="image/jpeg"), PROMPT],
                config=config,
            )
            return json.loads(r.text)
        except Exception as e:  # noqa: BLE001
            ultimo_error = e
            texto = str(e)
            # 429 = límite gratuito por minuto: esperar y reintentar
            if "429" in texto or "RESOURCE_EXHAUSTED" in texto or "503" in texto:
                time.sleep(20 * (intento + 1))
                continue
            raise
    raise ultimo_error


# ---------- tabla y Excel ----------
COLUMNAS = ["Hoja", "Fecha", "Sede", "Tarima", "Supervisor", "Caja", "EAN", "Descripción", "Piezas", "Nombre", "Estado", "Nota"]


def calcular_estado(df: pd.DataFrame) -> pd.Series:
    def fila(r):
        ok = ean_valido(r["EAN"]) and not str(r.get("Nota") or "").strip()
        return "✅" if ok else "⚠️ revisar"

    return df.apply(fila, axis=1)


def a_excel(df: pd.DataFrame, incluir_revision: bool) -> bytes:
    cols = ["Hoja", "Fecha", "Sede", "Tarima", "Supervisor",
            "Caja", "EAN", "Descripción", "Piezas", "Nombre"]
    if incluir_revision:
        cols += ["Estado", "Nota"]
    out = df[cols].copy()
    n = len(out)
    i_ean = cols.index("EAN")
    i_desc = cols.index("Descripción") + 1  # columnas de openpyxl empiezan en 1
    i_pz = cols.index("Piezas") + 1
    letra_pz = get_column_letter(i_pz)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        out.to_excel(w, index=False, sheet_name="Conteo")
        ws = w.sheets["Conteo"]

        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2E75B6")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        amarillo = PatternFill("solid", fgColor="FFFF00")
        estados = df["Estado"].tolist()
        for fila, estado in zip(ws.iter_rows(min_row=2, max_row=n + 1), estados):
            for c in fila:
                c.font = Font(name="Arial", size=10)
                if estado != "✅":  # fila por revisar -> sombreada de amarillo
                    c.fill = amarillo
            fila[i_ean].number_format = "@"  # EAN como texto

        ws.cell(n + 3, i_desc, "TOTAL DE PIEZAS").font = Font(name="Arial", bold=True)
        ws.cell(n + 3, i_pz, f"=SUM({letra_pz}2:{letra_pz}{n + 1})").font = Font(name="Arial", bold=True)

        for col in ws.columns:
            largo = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(largo + 3, 9), 60)
        ws.freeze_panes = "A2"
    return buf.getvalue()


# ---------- acceso con contraseña (opcional) ----------
def pedir_clave():
    """Si existe APP_PASSWORD en los secrets, pide la contraseña antes de mostrar la app."""
    try:
        clave = st.secrets["APP_PASSWORD"]
    except Exception:
        return  # sin contraseña configurada (uso local): no se pide nada
    if st.session_state.get("autorizado"):
        return
    ingreso = st.text_input("Contraseña de acceso", type="password")
    if ingreso:
        if hmac.compare_digest(ingreso, str(clave)):
            st.session_state["autorizado"] = True
            st.rerun()
        else:
            st.error("Contraseña incorrecta")
    st.stop()


# ---------- interfaz ----------
pedir_clave()
client = get_client()

archivos = st.file_uploader(
    "Sube las fotos de tus hojas (puedes elegir varias)",
    type=["jpg", "jpeg", "png"],
    accept_multiple_files=True,
)

if archivos:
    archivos = sorted(archivos, key=lambda f: f.name)
    filas, errores = [], []

    barra = st.progress(0.0, text="Leyendo hojas...")
    for i, f in enumerate(archivos):
        barra.progress(i / len(archivos), text=f"Leyendo {f.name} ({i + 1}/{len(archivos)})")
        try:
            data = leer_hoja(f.getvalue(), client)
        except Exception as e:  # noqa: BLE001
            errores.append(f"{f.name}: {e}")
            continue
        for r in data.get("filas", []):
            filas.append(
                {
                    "Hoja": f"Hoja {i + 1}",
                    "Fecha": data.get("fecha", ""),
                    "Sede": data.get("sede", ""),
                    "Tarima": data.get("tarima", ""),
                    "Supervisor": data.get("supervisor", ""),
                    "Caja": r.get("caja"),
                    "EAN": str(r.get("ean", "")).replace(" ", ""),
                    "Descripción": r.get("descripcion", ""),
                    "Piezas": r.get("piezas"),
                    "Nombre": r.get("nombre", ""),
                    "Nota": r.get("nota", "") if r.get("dudosa") else "",
                }
            )
    barra.empty()

    for e in errores:
        st.error(f"No pude leer {e}")

    with st.expander("Ver fotos subidas"):
        cols = st.columns(min(len(archivos), 4))
        for i, f in enumerate(archivos):
            cols[i % len(cols)].image(f.getvalue(), caption=f"Hoja {i + 1}: {f.name}")

    if filas:
        df = pd.DataFrame(filas)
        df["Caja"] = pd.to_numeric(df["Caja"], errors="coerce").astype("Int64")
        df["Piezas"] = pd.to_numeric(df["Piezas"], errors="coerce").astype("Int64")
        df["Estado"] = calcular_estado(df)
        df = df[COLUMNAS]

        clave = hashlib.md5(b"".join(f.getvalue() for f in archivos)).hexdigest()
        st.caption("Corrige lo que haga falta directamente en la tabla. Las filas ⚠️ son las que conviene revisar contra el papel.")
        editado = st.data_editor(
            df,
            num_rows="dynamic",
            use_container_width=True,
            key=f"editor_{clave}",
            column_config={
                "Estado": st.column_config.TextColumn(disabled=True),
                "Caja": st.column_config.NumberColumn(format="%d"),
                "Piezas": st.column_config.NumberColumn(format="%d"),
            },
        )
        editado["Estado"] = calcular_estado(editado)

        total = int(pd.to_numeric(editado["Piezas"], errors="coerce").fillna(0).sum())
        dudosas = int((editado["Estado"] != "✅").sum())
        c1, c2, c3 = st.columns(3)
        c1.metric("Filas", len(editado))
        c2.metric("Total de piezas", f"{total:,}")
        c3.metric("Por revisar", dudosas)

        incluir = st.checkbox("Incluir columnas de revisión (Estado y Nota) en el Excel (las filas por revisar siempre salen en amarillo)", value=False)
        st.download_button(
            "⬇️ DESCARGAR EXCEL",
            a_excel(editado, incluir),
            "inventario.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
        )
    elif not errores:
        st.warning("No encontré filas con datos. Prueba con una foto más recta y con buena luz.")
